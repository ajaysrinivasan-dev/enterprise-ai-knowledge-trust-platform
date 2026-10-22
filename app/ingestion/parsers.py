"""Safe parsers for supported local document formats."""

from io import BytesIO

from app.ingestion.models import ParsedDocument


def parse_document(filename: str, content: bytes, extension: str) -> ParsedDocument:
    """Extract text using an installed format adapter or report an actionable error."""

    if extension in {".txt", ".md", ".markdown"}:
        text = content.decode("utf-8-sig")
        page_count = None
    elif extension == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF support requires the PyMuPDF package") from exc
        document = fitz.open(stream=content, filetype="pdf")
        try:
            text = "\n".join(page.get_text() for page in document)
            page_count = document.page_count
        finally:
            document.close()
    elif extension == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires the python-docx package") from exc
        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        page_count = None
    else:
        raise ValueError(f"unsupported document type: {extension}")
    if not text.strip():
        raise ValueError("document contains no extractable text")
    return ParsedDocument(filename=filename, media_type=extension, text=text, page_count=page_count)
